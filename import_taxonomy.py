"""Convert the lab waste policy spreadsheet into taxonomy.json for Bin-Sight.

Run after editing the spreadsheet:   python import_taxonomy.py [data/lab_waste_taxonomy.xlsx]

The spreadsheet is the source of truth. Bin-Sight reads only taxonomy.json at runtime,
so the app itself doesn't need openpyxl.
"""
import json
import sys
from pathlib import Path

import openpyxl

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "data/lab_waste_taxonomy.xlsx")
OUT = Path("taxonomy.json")

# Policy categories that Bin-Sight must never assign to an ordinary bin.
OUT_OF_SCOPE = {"RAD", "ANI", "HICON", "UNK"}


def rows(ws):
    header = [str(h).strip() if h else "" for h in next(ws.iter_rows(max_row=1, values_only=True))]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0]:
            yield {h: ("" if v is None else str(v).strip()) for h, v in zip(header, r)}


def split(s, sep):
    return [x.strip() for x in s.split(sep) if x.strip()]


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)

    categories = {}
    for r in rows(wb["Categories"]):
        categories[r["Category ID"]] = {
            "id": r["Category ID"],
            "name": r["Name"],
            "container": r["Container (illustrative)"],
            "colour": "#" + r["Colour (hex)"].lstrip("#") if r["Colour (hex)"] else "#7F8C8D",
            "handling": r["Handling rules"],
            "notes": r["Escalation / notes"],
            "route_flag": r["Route flag"],
            "out_of_scope": r["Category ID"] in OUT_OF_SCOPE,
        }

    items = []
    for r in rows(wb["Items"]):
        items.append({
            "id": r["Item ID"],
            "name": r["Item"],
            "aliases": split(r["Search terms / aliases"], ","),
            "category_id": r["Category ID"],
            "steps": r["Pre-treatment / steps"],
            "flags": split(r["Hazard flags"], ";"),
            "do_not": r["Do NOT"],
            "why": r["Why (one line)"],
            "confidence": r["Confidence"],
        })

    rules = []
    for r in rows(wb["Incompatibilities"]):
        rules.append({
            "id": r["Rule ID"], "a": r["Flag A"], "b": r["Flag B"],
            "hazard": r["Hazard"], "rule": r["Rule for the app"], "severity": r["Severity"],
        })

    workflows = {}
    if "Workflows" in wb.sheetnames:
        for r in rows(wb["Workflows"]):
            workflows.setdefault(r["Workflow"], []).append({"stage": r["Stage"], "item_id": r["Item ID"]})

    # Sanity checks: every item points at a real category, and every category maps to a bin.
    bins = json.loads(Path("waste_categories.json").read_text())
    mapped = {t for b in bins for t in b.get("taxonomy", [])}
    bad_items = [i["id"] for i in items if i["category_id"] not in categories]
    unmapped = sorted(set(categories) - mapped - OUT_OF_SCOPE)

    OUT.write_text(json.dumps(
        {"source": SRC.name, "categories": categories, "items": items,
         "incompatibilities": rules, "workflows": workflows},
        indent=1, ensure_ascii=False))

    print(f"Wrote {OUT}: {len(categories)} categories, {len(items)} items, "
          f"{len(rules)} mixing rules, {len(workflows)} workflows")
    if bad_items:
        print(f"⚠️  Items with an unknown category: {', '.join(bad_items)}")
    if unmapped:
        print(f"⚠️  Policy categories not mapped to any bin in waste_categories.json: {', '.join(unmapped)}")
        print("   Add them to a bin's \"taxonomy\" list, or Bin-Sight will report them as needing a missing bin.")


if __name__ == "__main__":
    main()
